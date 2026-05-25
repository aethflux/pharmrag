# -*- coding: utf-8 -*-
"""
Single-turn attachment parsing for chat questions.

Attachments are temporary context. They are not written to the long-term
medical or personal knowledge bases unless the caller explicitly stores a
derived summary elsewhere.
"""

from __future__ import annotations

import base64
import csv
import json
import mimetypes
import re
from dataclasses import dataclass, field
from io import BytesIO, StringIO
from pathlib import Path
from typing import Callable

from config import APP_CONFIG
from rag.ocr import OcrUnavailableError, image_bytes_to_text, pdf_bytes_to_ocr_text

try:
    from docx import Document as DocxDocument
except ImportError:  # pragma: no cover
    DocxDocument = None

try:
    from pypdf import PdfReader
except ImportError:  # pragma: no cover
    PdfReader = None


CHAT_ATTACHMENT_EXTENSIONS = {
    ".txt",
    ".md",
    ".json",
    ".jsonl",
    ".csv",
    ".docx",
    ".pdf",
    ".png",
    ".jpg",
    ".jpeg",
    ".webp",
}

IMAGE_ATTACHMENT_EXTENSIONS = {".png", ".jpg", ".jpeg", ".webp"}
TEXT_ATTACHMENT_EXTENSIONS = {".txt", ".md", ".json", ".jsonl", ".csv"}

VisionSummarizer = Callable[[str, str, str], str]


@dataclass
class AttachmentContext:
    filename: str
    file_type: str
    size: int
    content_type: str = ""
    extracted_text: str = ""
    ocr_text: str = ""
    vision_summary: str = ""
    warnings: list[str] = field(default_factory=list)
    can_reference: bool = False
    is_image: bool = False
    data_url: str = ""

    def public_report(self) -> dict[str, object]:
        return {
            "filename": self.filename,
            "file_type": self.file_type,
            "size": self.size,
            "content_type": self.content_type,
            "has_text": bool(self.extracted_text.strip() or self.ocr_text.strip()),
            "has_vision_summary": bool(self.vision_summary.strip()),
            "warnings": self.warnings,
            "can_reference": self.can_reference,
            "is_image": self.is_image,
        }

    def reference_text(self) -> str:
        sections: list[str] = []
        if self.extracted_text.strip():
            sections.append("文本抽取：\n" + _limit_text(self.extracted_text))
        if self.ocr_text.strip():
            sections.append("OCR 识别：\n" + _limit_text(self.ocr_text))
        if self.vision_summary.strip():
            sections.append("视觉摘要：\n" + _limit_text(self.vision_summary))
        if self.warnings:
            sections.append("处理提示：\n" + "\n".join(f"- {item}" for item in self.warnings))
        return "\n\n".join(sections).strip()


def _limit_text(value: str) -> str:
    text = re.sub(r"\n{3,}", "\n\n", str(value or "")).strip()
    limit = max(APP_CONFIG["attachment_preview_chars"], 1000)
    if len(text) <= limit:
        return text
    return text[:limit].rstrip() + "\n...[内容已截断]"


def _decode_text(data: bytes) -> str:
    for encoding in ("utf-8-sig", "utf-8", "gb18030"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors="ignore")


def _read_text_attachment(filename: str, data: bytes) -> str:
    suffix = Path(filename).suffix.lower()
    decoded = _decode_text(data).strip()
    if suffix == ".json":
        try:
            return json.dumps(json.loads(decoded), ensure_ascii=False, indent=2)
        except Exception:
            return decoded
    if suffix == ".csv":
        try:
            reader = csv.reader(StringIO(decoded))
            rows = [", ".join(cell.strip() for cell in row) for row in reader]
            return "\n".join(row for row in rows if row.strip())
        except Exception:
            return decoded
    return decoded


def _read_docx_attachment(data: bytes) -> str:
    if DocxDocument is None:
        raise ValueError("缺少 python-docx 依赖，无法解析 docx。")
    document = DocxDocument(BytesIO(data))
    paragraphs = [paragraph.text.strip() for paragraph in document.paragraphs if paragraph.text.strip()]
    table_rows: list[str] = []
    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if cells:
                table_rows.append(" | ".join(cells))
    return "\n".join(paragraphs + table_rows).strip()


def _read_pdf_attachment(data: bytes) -> tuple[str, str, list[str]]:
    warnings: list[str] = []
    if PdfReader is None:
        raise ValueError("缺少 pypdf 依赖，无法解析 PDF。")

    reader = PdfReader(BytesIO(data))
    pages: list[str] = []
    for index, page in enumerate(reader.pages, start=1):
        text = (page.extract_text() or "").strip()
        if text:
            pages.append(f"[page {index}]\n{text}")

    extracted = "\n\n".join(pages).strip()
    if extracted:
        return extracted, "", warnings

    try:
        return "", pdf_bytes_to_ocr_text(data), warnings
    except OcrUnavailableError as exc:
        warnings.append(f"PDF 未抽取到文本，OCR 不可用：{exc}")
        return "", "", warnings


def _mime_type(filename: str, content_type: str = "") -> str:
    if content_type:
        return content_type
    guessed, _ = mimetypes.guess_type(filename)
    return guessed or "application/octet-stream"


def _to_data_url(filename: str, data: bytes, content_type: str = "") -> str:
    mime_type = _mime_type(filename, content_type)
    encoded = base64.b64encode(data).decode("ascii")
    return f"data:{mime_type};base64,{encoded}"


def parse_attachment_bytes(
    filename: str,
    data: bytes,
    *,
    content_type: str = "",
    question: str = "",
    vision_summarizer: VisionSummarizer | None = None,
) -> AttachmentContext:
    suffix = Path(filename).suffix.lower()
    context = AttachmentContext(
        filename=Path(filename).name or "attachment",
        file_type=suffix,
        size=len(data),
        content_type=_mime_type(filename, content_type),
        is_image=suffix in IMAGE_ATTACHMENT_EXTENSIONS,
    )

    if not data:
        context.warnings.append("文件内容为空。")
        return context
    if suffix not in CHAT_ATTACHMENT_EXTENSIONS:
        supported = ", ".join(sorted(CHAT_ATTACHMENT_EXTENSIONS))
        context.warnings.append(f"不支持的附件类型：{suffix or '无扩展名'}。支持：{supported}")
        return context
    if len(data) > APP_CONFIG["attachment_max_bytes"]:
        context.warnings.append(f"附件超过大小限制：{APP_CONFIG['attachment_max_bytes']} bytes。")
        return context

    try:
        if suffix in TEXT_ATTACHMENT_EXTENSIONS:
            context.extracted_text = _read_text_attachment(filename, data)
        elif suffix == ".docx":
            context.extracted_text = _read_docx_attachment(data)
        elif suffix == ".pdf":
            context.extracted_text, context.ocr_text, pdf_warnings = _read_pdf_attachment(data)
            context.warnings.extend(pdf_warnings)
        elif suffix in IMAGE_ATTACHMENT_EXTENSIONS:
            context.data_url = _to_data_url(filename, data, content_type)
            try:
                context.ocr_text = image_bytes_to_text(data)
            except OcrUnavailableError as exc:
                context.warnings.append(f"图片 OCR 不可用：{exc}")
            if vision_summarizer is not None:
                try:
                    context.vision_summary = vision_summarizer(context.data_url, question, context.filename)
                except Exception as exc:
                    context.warnings.append(f"视觉模型摘要失败：{exc}")
    except Exception as exc:
        context.warnings.append(f"附件解析失败：{exc}")

    context.can_reference = bool(
        context.extracted_text.strip()
        or context.ocr_text.strip()
        or context.vision_summary.strip()
    )
    return context


def build_attachments_prompt(attachments: list[AttachmentContext]) -> str:
    usable = [item for item in attachments if item.can_reference or item.warnings]
    if not usable:
        return ""

    sections = []
    for index, item in enumerate(usable, start=1):
        body = item.reference_text()
        if not body:
            body = "未能提取出可引用内容。"
        sections.append(
            f"【附件 {index}：{item.filename}｜{item.file_type or 'unknown'}｜{item.size} bytes】\n{body}"
        )
    return "\n\n".join(sections).strip()


def build_personal_knowledge_attachment_summary(
    prompt: str,
    attachments: list[AttachmentContext],
) -> str:
    attachment_text = build_attachments_prompt(attachments)
    sections = [
        "## 用户选择保存的单次提问",
        "",
        "以下内容来自用户本次提问，用户明确选择加入个人信息库：",
        "",
        prompt.strip(),
    ]
    if attachment_text:
        sections.extend(
            [
                "",
                "## 本轮附件可解析摘要",
                "",
                attachment_text,
            ]
        )
    return "\n".join(sections).strip()


def has_medical_visual_attachment(question: str, attachments: list[AttachmentContext]) -> bool:
    if not any(item.is_image for item in attachments):
        return False
    text = question.lower()
    markers = [
        "外伤",
        "伤口",
        "出血",
        "皮疹",
        "红肿",
        "感染",
        "影像",
        "ct",
        "x光",
        "核磁",
        "片子",
        "报告",
        "检查",
        "化验",
        "药盒",
        "说明书",
        "药品",
    ]
    return any(marker in text for marker in markers)
